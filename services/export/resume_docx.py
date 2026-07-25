import io

from docx import Document

FIXED_SECTION_LABELS = {
    "summary": "Summary",
    "experience": "Experience",
    "education": "Education",
    "skills": "Skills",
    "projects": "Projects",
    "certifications": "Certifications",
    "awards": "Awards",
    "volunteer": "Volunteer",
    "publications": "Publications",
    "interests": "Interests",
    "references": "References",
}


def get_section_label(key: str, custom_sections: list) -> str:
    if key.startswith("custom:"):
        section_id = key[len("custom:") :]
        section = next((s for s in custom_sections if s.get("id") == section_id), None)
        return section.get("title", "Custom Section") if section else "Custom Section"
    return FIXED_SECTION_LABELS.get(key, key)


def contact_line(resume: dict) -> str:
    personal_info = resume.get("personalInfo") or {}
    parts = [personal_info.get(k) for k in ("email", "phone", "address", "linkedin", "github", "portfolio")]
    return "  |  ".join(p for p in parts if p)


def build_section_lines(resume: dict, key: str) -> list:
    """Returns a list of {"text", "bold", "bullet"} dicts for the section, or [] if empty."""
    lines = []

    if key == "summary":
        if resume.get("summary"):
            lines.append({"text": resume["summary"]})
    elif key == "experience":
        # The Builder pre-seeds one blank row so the editor isn't empty on
        # first load — skip it here until the person fills something in.
        for exp in resume.get("experience", []):
            if not (exp.get("company") or exp.get("position")):
                continue
            lines.append({"text": f"{exp.get('position', '')}, {exp.get('company', '')}", "bold": True})
            end = "Present" if exp.get("isCurrent") else (exp.get("endDate") or "")
            lines.append({"text": f"{exp.get('startDate', '')} - {end}"})
            for achievement in exp.get("achievements", []):
                if achievement:
                    lines.append({"text": achievement, "bullet": True})
    elif key == "education":
        for edu in resume.get("education", []):
            if not (edu.get("school") or edu.get("degree") or edu.get("field")):
                continue
            lines.append(
                {"text": f"{edu.get('degree', '')} in {edu.get('field', '')}, {edu.get('school', '')} ({edu.get('startDate', '')} - {edu.get('endDate') or ''})"}
            )
    elif key == "skills":
        skills = resume.get("skills", [])
        if skills:
            lines.append({"text": ", ".join(s.get("name", "") for s in skills)})
    elif key == "projects":
        for project in resume.get("projects", []):
            lines.append({"text": project.get("title", ""), "bold": True})
            lines.append({"text": project.get("description", "")})
    elif key == "certifications":
        for cert in resume.get("certifications", []):
            lines.append({"text": f"{cert.get('name', '')}, {cert.get('issuer', '')} ({cert.get('issueDate', '')})"})
    elif key == "awards":
        for award in resume.get("awards", []):
            lines.append({"text": f"{award.get('title', '')}, {award.get('issuer', '')} ({award.get('date', '')})"})
    elif key == "volunteer":
        for entry in resume.get("volunteer", []):
            end = entry.get("endDate") or "Present"
            lines.append({"text": f"{entry.get('role', '')}, {entry.get('organization', '')} ({entry.get('startDate', '')} - {end})"})
            if entry.get("description"):
                lines.append({"text": entry["description"]})
    elif key == "publications":
        for pub in resume.get("publications", []):
            lines.append({"text": f"{pub.get('title', '')}, {pub.get('publisher', '')} ({pub.get('date', '')})"})
    elif key == "interests":
        interests = [i for i in resume.get("interests", []) if i]
        if interests:
            lines.append({"text": ", ".join(interests)})
    elif key == "references":
        for ref in resume.get("references", []):
            lines.append({"text": f"{ref.get('name', '')}, {ref.get('relationship', '')} ({ref.get('contactInfo', '')})"})
    elif key.startswith("custom:"):
        section_id = key[len("custom:") :]
        section = next((s for s in resume.get("customSections", []) if s.get("id") == section_id), None)
        if section:
            for line in section.get("content", "").split("\n"):
                lines.append({"text": line})

    return lines


def build_resume_docx(resume: dict) -> io.BytesIO:
    document = Document()
    personal_info = resume.get("personalInfo") or {}

    name_paragraph = document.add_paragraph()
    name_run = name_paragraph.add_run(personal_info.get("fullName") or "Untitled")
    name_run.bold = True

    document.add_paragraph(personal_info.get("jobTitle", ""))
    document.add_paragraph(contact_line(resume))

    for key in resume.get("sectionOrder", []):
        lines = build_section_lines(resume, key)
        if not lines:
            continue
        document.add_heading(get_section_label(key, resume.get("customSections", [])), level=2)
        for line in lines:
            paragraph = document.add_paragraph(style="List Bullet" if line.get("bullet") else None)
            run = paragraph.add_run(line["text"])
            run.bold = bool(line.get("bold"))

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
